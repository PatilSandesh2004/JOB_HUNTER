"""Resume text extraction (PDF/DOCX/TXT) and structured profile parsing.

Heuristics always run; when an LLM is configured its structured extraction is merged on top.
Anything that cannot be found is left empty rather than invented.
"""

import io
import logging
import re
from datetime import date

from pydantic import ValidationError

from ai_service.app.core.errors import LLMUnavailableError, ResumeParseError
from ai_service.app.integrations.llm.llm_client import LLMClient
from ai_service.app.schemas.candidate import CandidateProfile, EducationItem, WorkExperience
from ai_service.app.services.skills.catalog import extract_skills, normalize_skill_list

logger = logging.getLogger("jobpilot.resume")

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE_RE = re.compile(r"(?<!\d)(\+?\d{1,3}[\s.-]?)?(\(?\d{2,5}\)?[\s.-]?)\d{3,5}[\s.-]?\d{3,5}(?!\d)")
_LINKEDIN_RE = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w%-]+/?", re.I)
_GITHUB_RE = re.compile(r"(?:https?://)?github\.com/[\w-]+/?", re.I)
_YEARS_RE = re.compile(
    r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:professional\s+|industry\s+|work\s+)?experience", re.I
)
_MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
_DATE = rf"(?:(?:{_MONTHS})[a-z]*\.?\s+|\d{{1,2}}/)?(?:19|20)\d{{2}}"
_RANGE_RE = re.compile(rf"({_DATE})\s*(?:-|–|—|to)\s*({_DATE}|present|current|now|till date|ongoing)", re.I)

_LLM_SYSTEM = (
    "You extract structured data from resumes. Respond with a single JSON object only. "
    "Use null or empty lists for anything not explicitly present. Never guess or invent values."
)
_LLM_USER = """Extract this resume into JSON with keys:
name, email, phone, location, current_role, current_company, summary (<=2 sentences, from the resume text),
years_of_experience (number, total professional experience), linkedin_url, github_url, portfolio_url,
skills (list of strings),
work_experience (list of {{company, title, location, start_date, end_date, description}}),
education (list of {{institution, degree, field_of_study, graduation_year}}).

Resume:
\"\"\"
{text}
\"\"\""""


class ResumeParserService:
    def __init__(self, llm: LLMClient | None = None) -> None:
        self.llm = llm

    # ---- text extraction -------------------------------------------------
    def extract_text(self, data: bytes, filename: str) -> str:
        name = filename.lower()
        try:
            if name.endswith(".pdf"):
                text = self._pdf_text(data)
            elif name.endswith(".docx"):
                text = self._docx_text(data)
            elif name.endswith(".txt"):
                text = data.decode("utf-8", errors="ignore")
            else:
                raise ResumeParseError(f"Unsupported file type. Use one of: {', '.join(SUPPORTED_EXTENSIONS)}")
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError(f"Could not read resume: {exc}") from exc
        if len(text.strip()) < 30:
            raise ResumeParseError("No readable text found (is the PDF a scanned image?)")
        return text

    @staticmethod
    def _pdf_text(data: bytes) -> str:
        try:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(data)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as exc:
            logger.info("pdfplumber failed (%s); falling back to pypdf", exc)
            from pypdf import PdfReader

            return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)

    @staticmethod
    def _docx_text(data: bytes) -> str:
        from docx import Document

        document = Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    # ---- parsing ---------------------------------------------------------
    async def parse(self, text: str) -> CandidateProfile:
        profile = self.parse_heuristic(text)
        if self.llm is None or not self.llm.available:
            return profile
        try:
            # A full resume as JSON (experience + education) needs a generous output budget.
            data = await self.llm.complete_json(
                _LLM_SYSTEM, _LLM_USER.format(text=text[:12000]), temperature=0.0, max_tokens=6000
            )
            return _merge(profile, data)
        except (LLMUnavailableError, ValueError) as exc:
            logger.warning("LLM resume extraction unavailable, using heuristics only: %s", exc)
            return profile

    def parse_heuristic(self, text: str) -> CandidateProfile:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        email = _first(_EMAIL_RE, text)
        return CandidateProfile(
            name=_guess_name(lines),
            email=email,
            phone=_guess_phone(text),
            linkedin_url=_first(_LINKEDIN_RE, text),
            github_url=_first(_GITHUB_RE, text),
            years_of_experience=_years_of_experience(text),
            skills=extract_skills(text),
        )


def _first(pattern: re.Pattern[str], text: str) -> str | None:
    match = pattern.search(text)
    return match.group(0).strip() if match else None


def _guess_phone(text: str) -> str | None:
    for match in _PHONE_RE.finditer(text):
        candidate = match.group(0).strip()
        digits = re.sub(r"\D", "", candidate)
        if 10 <= len(digits) <= 15 and not re.fullmatch(r"(19|20)\d{2}\D*(19|20)\d{2}", candidate):
            return candidate
    return None


def _guess_name(lines: list[str]) -> str:
    for line in lines[:6]:
        cleaned = re.sub(r"\s+", " ", line)
        if (
            2 <= len(cleaned.split()) <= 4
            and len(cleaned) <= 50
            and re.fullmatch(r"[A-Za-z][A-Za-z .'-]+", cleaned)
            and cleaned.lower() not in {"curriculum vitae", "resume"}
        ):
            return cleaned.title() if cleaned.isupper() else cleaned
    return ""


def _years_of_experience(text: str) -> float:
    explicit = [float(m.group(1)) for m in _YEARS_RE.finditer(text)]
    if explicit:
        return max(explicit)
    return _years_from_date_ranges(text)


def _years_from_date_ranges(text: str) -> float:
    """Sum non-overlapping employment ranges like 'Jan 2020 - Present'."""
    intervals: list[tuple[int, int]] = []
    for match in _RANGE_RE.finditer(text):
        start, end = _month_index(match.group(1)), _month_index(match.group(2))
        if start is not None and end is not None and end > start:
            intervals.append((start, end))
    if not intervals:
        return 0.0
    intervals.sort()
    total, (cur_start, cur_end) = 0, intervals[0]
    for start, end in intervals[1:]:
        if start <= cur_end:
            cur_end = max(cur_end, end)
        else:
            total += cur_end - cur_start
            cur_start, cur_end = start, end
    total += cur_end - cur_start
    return round(min(total / 12, 50.0), 1)


def _month_index(value: str) -> int | None:
    value = value.strip().lower()
    if value in {"present", "current", "now", "till date", "ongoing"}:
        today = date.today()
        return today.year * 12 + today.month
    year_match = re.search(r"(19|20)\d{2}", value)
    if not year_match:
        return None
    month = 1
    month_match = re.match(rf"({_MONTHS})", value)
    if month_match:
        month = [m[:3] for m in _MONTHS.split("|")].index(month_match.group(1)[:3]) + 1
    elif slash := re.match(r"(\d{1,2})/", value):
        month = max(1, min(12, int(slash.group(1))))
    return int(year_match.group(0)) * 12 + month


def _merge(base: CandidateProfile, llm_data: dict) -> CandidateProfile:
    """Overlay LLM values on heuristic results.

    Contact details matched by regex are exact, so they win and the LLM only fills blanks.
    Descriptive fields need language understanding, so the LLM wins when it returns a value.
    """
    merged = base.model_dump()

    def llm_str(key: str) -> str | None:
        value = llm_data.get(key)
        return value.strip() if isinstance(value, str) and value.strip() else None

    for key in ("email", "phone", "linkedin_url", "github_url"):
        merged[key] = merged.get(key) or llm_str(key)
    for key in ("name", "location", "current_role", "current_company", "summary", "portfolio_url"):
        merged[key] = llm_str(key) or merged.get(key)
    if merged["name"] and merged["name"].isupper():
        merged["name"] = merged["name"].title()  # "SANDESH PATIL" -> "Sandesh Patil"

    years = llm_data.get("years_of_experience")
    if isinstance(years, int | float) and 0 < years <= 50:
        merged["years_of_experience"] = float(years)

    llm_skills = [s for s in llm_data.get("skills") or [] if isinstance(s, str)]
    merged["skills"] = normalize_skill_list([*merged["skills"], *llm_skills])[:60]
    merged["work_experience"] = _valid_items(llm_data.get("work_experience"), WorkExperience)
    merged["education"] = _valid_items(llm_data.get("education"), EducationItem)
    try:
        return CandidateProfile.model_validate(merged)
    except ValidationError:
        merged["email"] = base.email
        return CandidateProfile.model_validate(merged)


def _valid_items(items: object, model: type) -> list:
    if not isinstance(items, list):
        return []
    valid = []
    for item in items:
        try:
            valid.append(model.model_validate(item).model_dump())
        except ValidationError:
            continue
    return valid
